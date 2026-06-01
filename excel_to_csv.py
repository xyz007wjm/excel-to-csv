import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import os
import csv
import webbrowser

# 根据文件格式选择对应的库
_HAS_OPENPYXL = False
_HAS_XLRD = False
try:
    import openpyxl
    _HAS_OPENPYXL = True
except ImportError:
    pass
try:
    import xlrd
    _HAS_XLRD = True
except ImportError:
    pass

# 软件信息
APP_NAME = "Excel 转 CSV 工具"
APP_VERSION = "1.0"
APP_AUTHOR = "Jack Wang"
APP_PHONE = "15657317797"
APP_BLOG = "https://wangjinming.com"


class ExcelToCsvApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_NAME} v{APP_VERSION}")
        self.root.geometry("700x550")
        self.root.minsize(600, 450)

        self.files = []

        self._setup_ui()

    def _setup_ui(self):
        # 标题
        title = tk.Label(self.root, text=APP_NAME, font=("Arial", 16, "bold"))
        title.pack(pady=(15, 2))

        version_label = tk.Label(
            self.root,
            text=f"版本 {APP_VERSION}",
            font=("Arial", 9),
            fg="gray",
        )
        version_label.pack(pady=(0, 5))

        desc = tk.Label(
            self.root,
            text="支持 .xlsx / .xls 文件，分隔符为分号 (;)，输出到 output 文件夹",
            font=("Arial", 10),
            fg="gray",
        )
        desc.pack(pady=(0, 10))

        # 按钮区域
        btn_frame = tk.Frame(self.root)
        btn_frame.pack(fill=tk.X, padx=20, pady=(0, 10))

        add_btn = tk.Button(
            btn_frame,
            text="添加 Excel 文件",
            command=self._add_files,
            bg="#4CAF50",
            fg="white",
            padx=15,
            pady=5,
            font=("Arial", 10),
        )
        add_btn.pack(side=tk.LEFT, padx=(0, 10))

        clear_btn = tk.Button(
            btn_frame,
            text="清空列表",
            command=self._clear_files,
            bg="#f44336",
            fg="white",
            padx=15,
            pady=5,
            font=("Arial", 10),
        )
        clear_btn.pack(side=tk.LEFT)

        # 文件列表
        list_frame = tk.Frame(self.root)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 10))

        list_label = tk.Label(list_frame, text="待转换文件列表：", anchor=tk.W, font=("Arial", 10))
        list_label.pack(fill=tk.X)

        columns = ("filename", "sheets", "path")
        self.tree = ttk.Treeview(list_frame, columns=columns, show="headings", height=10)
        self.tree.heading("filename", text="文件名")
        self.tree.heading("sheets", text="工作表数")
        self.tree.heading("path", text="路径")
        self.tree.column("filename", width=250)
        self.tree.column("sheets", width=80, anchor=tk.CENTER)
        self.tree.column("path", width=300)
        self.tree.pack(fill=tk.BOTH, expand=True, side=tk.LEFT)

        scrollbar = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.tree.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.configure(yscrollcommand=scrollbar.set)

        # 右键菜单
        self.popup = tk.Menu(self.root, tearoff=0)
        self.popup.add_command(label="移除选中文件", command=self._remove_selected)
        self.tree.bind("<Button-3>", self._show_popup)

        # 删除键删除
        self.root.bind("<Delete>", lambda e: self._remove_selected())

        # 转换按钮
        self.convert_btn = tk.Button(
            self.root,
            text="开始转换",
            command=self._convert,
            bg="#2196F3",
            fg="white",
            padx=30,
            pady=8,
            font=("Arial", 12, "bold"),
            state=tk.DISABLED,
        )
        self.convert_btn.pack(pady=(0, 10))

        # 版权信息栏
        copyright_frame = tk.Frame(self.root, bd=1, relief=tk.SUNKEN)
        copyright_frame.pack(fill=tk.X, side=tk.BOTTOM)

        copyright_text = f"\u7248\u6743\u6240\u6709 \u00a9 {APP_AUTHOR}  |  \u7535\u8bdd\uff1a{APP_PHONE}  |  "
        copyright_label = tk.Label(
            copyright_frame,
            text=copyright_text,
            anchor=tk.W,
            padx=10,
            pady=3,
            font=("Arial", 9),
        )
        copyright_label.pack(side=tk.LEFT)

        blog_link = tk.Label(
            copyright_frame,
            text=APP_BLOG,
            fg="#2196F3",
            cursor="hand2",
            font=("Arial", 9, "underline"),
            padx=0,
            pady=3,
        )
        blog_link.pack(side=tk.LEFT)
        blog_link.bind("<Button-1>", lambda e: webbrowser.open(APP_BLOG))

        # 状态栏
        self.status_var = tk.StringVar(value="就绪")
        status_bar = tk.Label(
            self.root,
            textvariable=self.status_var,
            bd=1,
            relief=tk.SUNKEN,
            anchor=tk.W,
            padx=10,
        )
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)

    def _add_files(self):
        paths = filedialog.askopenfilenames(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx *.xls"), ("所有文件", "*.*")],
        )
        if not paths:
            return

        for path in paths:
            if path in self.files:
                continue

            ext = os.path.splitext(path)[1].lower()
            sheet_count = 0
            ok = False

            try:
                if ext == ".xlsx":
                    if not _HAS_OPENPYXL:
                        raise RuntimeError("缺少 openpyxl 库，无法读取 .xlsx 文件")
                    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                    sheet_count = len(wb.sheetnames)
                    wb.close()
                    ok = True
                elif ext == ".xls":
                    if not _HAS_XLRD:
                        raise RuntimeError("缺少 xlrd 库，无法读取 .xls 文件")
                    wb = xlrd.open_workbook(path)
                    sheet_count = wb.nsheets
                    wb.release_resources()
                    ok = True
                else:
                    raise RuntimeError(f"不支持的文件格式：{ext}")
            except Exception as e:
                messagebox.showerror("读取失败", f"无法读取文件：{os.path.basename(path)}\n{str(e)}")

            if ok:
                self.files.append(path)
                self.tree.insert(
                    "",
                    tk.END,
                    values=(
                        os.path.basename(path),
                        sheet_count,
                        os.path.dirname(path),
                    ),
                )

        self._update_status()
        self.convert_btn.config(state=tk.NORMAL if self.files else tk.DISABLED)

    def _clear_files(self):
        if not self.files:
            return
        if messagebox.askyesno("确认清空", "确定要清空所有文件吗？"):
            self.files.clear()
            self.tree.delete(*self.tree.get_children())
            self.convert_btn.config(state=tk.DISABLED)
            self.status_var.set("就绪")

    def _remove_selected(self):
        selected = self.tree.selection()
        if not selected:
            return
        for item in selected:
            values = self.tree.item(item, "values")
            path = os.path.join(values[2], values[0])
            if path in self.files:
                self.files.remove(path)
            self.tree.delete(item)
        self._update_status()
        self.convert_btn.config(state=tk.NORMAL if self.files else tk.DISABLED)

    def _show_popup(self, event):
        item = self.tree.identify_row(event.y)
        if item:
            self.tree.selection_set(item)
            self.popup.post(event.x_root, event.y_root)

    def _update_status(self):
        count = len(self.files)
        self.status_var.set(f"共 {count} 个文件待转换" if count else "就绪")

    def _convert_single(self, file_path):
        base = os.path.splitext(os.path.basename(file_path))[0]
        out_dir = os.path.join(os.path.dirname(file_path), "output")
        os.makedirs(out_dir, exist_ok=True)
        out_name = f"{base}.csv"
        out_path = os.path.join(out_dir, out_name)

        ext = os.path.splitext(file_path)[1].lower()

        try:
            if ext == ".xlsx":
                rows = self._read_xlsx(file_path)
            elif ext == ".xls":
                rows = self._read_xls(file_path)
            else:
                return False, f"不支持的文件格式：{ext}"

            with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f, delimiter=";")
                for row in rows:
                    writer.writerow(["" if v is None else str(v) for v in row])

            return True, out_name
        except Exception as e:
            return False, str(e)

    @staticmethod
    def _read_xlsx(file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        wb.close()
        return rows

    @staticmethod
    def _read_xls(file_path):
        wb = xlrd.open_workbook(file_path)
        sheet = wb.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            rows.append(sheet.row_values(r))
        wb.release_resources()
        return rows

    def _convert(self):
        if not self.files:
            return

        success_count = 0
        fail_count = 0
        errors = []

        self.convert_btn.config(state=tk.DISABLED, text="转换中...")
        self.root.update()

        for i, file_path in enumerate(self.files):
            self.status_var.set(f"正在转换 ({i+1}/{len(self.files)})：{os.path.basename(file_path)}")
            self.root.update()

            ok, result = self._convert_single(file_path)
            if ok:
                success_count += 1
            else:
                fail_count += 1
                errors.append(f"{os.path.basename(file_path)}：{result}")

        self.convert_btn.config(state=tk.NORMAL, text="开始转换")

        msg = f"转换完成！\n成功：{success_count} 个"
        if fail_count:
            msg += f"\n失败：{fail_count} 个"
            msg += "\n\n" + "\n".join(errors)

        messagebox.showinfo("转换结果", msg)
        self.status_var.set(f"完成：成功 {success_count} 个" + (f"，失败 {fail_count} 个" if fail_count else ""))


def main():
    root = tk.Tk()
    app = ExcelToCsvApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
